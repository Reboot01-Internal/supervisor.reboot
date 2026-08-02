from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/reemhalwachi/Desktop/reboot/super")
IMG = ROOT / "tmp/taskflow_proposal/source_images"
OUT = ROOT / "TaskFlow_Project_Proposal.docx"

NAVY = "15243A"
TEAL = "159F97"
PURPLE = "7567E8"
INK = "243247"
MUTED = "697588"
PALE = "F1F6F6"
WHITE = "FFFFFF"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=140, bottom=110, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def font_run(run, size=11, bold=False, color=INK, italic=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def add_text(doc, text, size=11, bold=False, color=INK, align=None,
             before=0, after=7, line=1.18, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    font_run(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_picture(doc, filename, width=6.1, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(IMG / filename), width=Inches(width))
    if caption:
        c = add_text(
            doc, caption, size=9, color=MUTED,
            align=WD_ALIGN_PARAGRAPH.CENTER, after=9, italic=True
        )
        c.paragraph_format.keep_with_next = True


def add_feature_table(doc, items):
    table = doc.add_table(rows=1, cols=len(items))
    table.autofit = False
    usable = 9360
    widths = [usable // len(items)] * len(items)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(usable))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    remove_table_borders(table)
    for idx, (title, text) in enumerate(items):
        cell = table.cell(0, idx)
        cell.width = Inches(6.5 / len(items))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, PALE if idx % 2 == 0 else "F6F4FF")
        set_cell_margins(cell, 150, 150, 150, 150)
        tcw = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
        tcw.set(qn("w:w"), str(widths[idx]))
        tcw.set(qn("w:type"), "dxa")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        font_run(p.add_run(title), size=11, bold=True, color=TEAL if idx % 2 == 0 else PURPLE)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.12
        font_run(p2.add_run(text), size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_process_step(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.05)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    font_run(p.add_run(f"{number}  "), size=11, bold=True, color=PURPLE)
    font_run(p.add_run(f"{title} — "), size=11, bold=True, color=NAVY)
    font_run(p.add_run(text), size=10.5, color=INK)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.38)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = rgb(INK)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.18

for name, size, color, before, after in (
    ("Title", 29, NAVY, 0, 7),
    ("Subtitle", 14, MUTED, 0, 15),
    ("Heading 1", 17, NAVY, 14, 8),
    ("Heading 2", 13, TEAL, 9, 5),
    ("Heading 3", 11, PURPLE, 7, 4),
):
    style = styles[name]
    style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = name != "Subtitle"
    style.font.color.rgb = rgb(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font_run(hp.add_run("TASKFLOW  |  PROJECT PROPOSAL"), size=8.5, bold=True, color=MUTED)

footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font_run(fp.add_run("Prepared by Reem Alhalwachi  •  July 2026"), size=8.5, color=MUTED)

# Page 1 — proposal centerpiece
add_text(doc, "PRODUCT PROPOSAL", size=10, bold=True, color=TEAL,
         align=WD_ALIGN_PARAGRAPH.CENTER, before=14, after=10)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(7)
font_run(p.add_run("TaskFlow"), size=29, bold=True, color=NAVY, name="Aptos Display")
add_text(doc, "A simpler way to guide students, manage work, and see progress",
         size=15, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
add_text(doc, "Prepared for a prospective company partner",
         size=10.5, bold=True, color=PURPLE, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)

add_picture(doc, "p1497_rId93.png", width=5.95,
            caption="A single dashboard gives management a clear view of students, supervisors, projects, tasks, and overall progress.")

add_heading(doc, "Proposal at a glance", 1)
add_text(
    doc,
    "TaskFlow is a ready-built web product for organisations that supervise learners or junior team members through project work. "
    "It replaces scattered messages and manual follow-up with one organised workspace for assignments, tasks, meetings, reminders, and progress reporting.",
    size=10.7, after=8
)
add_feature_table(doc, [
    ("Purpose", "Make supervision clear, consistent, and easier to manage."),
    ("Users", "Administrators, supervisors, and students or trainees."),
    ("Opportunity", "Pilot the product, adapt it to company needs, and prepare it for wider use."),
])

doc.add_page_break()

# Page 2 — problem and solution
add_heading(doc, "The business need", 1)
add_text(
    doc,
    "When supervision depends on spreadsheets, direct messages, and memory, it becomes difficult to know who needs support, "
    "what work is late, and whether follow-up happened. TaskFlow creates one shared process without making the experience complicated.",
    size=11, after=9
)
add_feature_table(doc, [
    ("Less manual coordination", "Assignments, project boards, meetings, and reminders are managed in one place."),
    ("Better visibility", "Managers and supervisors can quickly see workload, progress, overdue work, and activity."),
    ("More accountability", "Every learner knows their tasks, deadlines, supervisor, and next action."),
])

add_heading(doc, "How the product works", 1)
add_process_step(doc, "1", "Assign people", "Administrators connect each student or trainee with the right supervisor.")
add_process_step(doc, "2", "Create a workspace", "A project board gives the team one organised place for the work.")
add_process_step(doc, "3", "Guide the work", "Supervisors add tasks, deadlines, priorities, notes, and checklists.")
add_process_step(doc, "4", "Stay connected", "Meetings, notifications, and Discord reminders keep everyone informed.")
add_process_step(doc, "5", "Review progress", "Dashboards and reports show what is completed and where support is needed.")

add_picture(doc, "p1524_rId99.png", width=5.8,
            caption="Administrators can see available students and assign them to supervisors through a clear visual workflow.")

doc.add_page_break()

# Page 3 — core product experience
add_heading(doc, "A clear experience for everyday work", 1)
add_text(
    doc,
    "The interface uses familiar boards, lists, and task cards, so users can understand the workflow quickly. "
    "Supervisors manage the plan while students focus on their assigned work.",
    size=11, after=6
)
add_picture(doc, "p1649_rId127.png", width=6.05,
            caption="Supervisors organise project work into simple stages and can see task status at a glance.")

add_heading(doc, "Each task carries the full context", 2)
add_text(
    doc,
    "A task can include a due date, priority, description, assignee, comments, labels, and a checklist. "
    "This reduces repeated questions and gives the learner a clear definition of what needs to be done.",
    size=10.7, after=5
)
add_picture(doc, "p1575_rId111.png", width=5.6,
            caption="Task details keep instructions, deadlines, ownership, and smaller action items together.")

doc.add_page_break()

# Page 4 — communication and value
add_heading(doc, "Supervision continues beyond the task board", 1)
add_text(
    doc,
    "TaskFlow brings meetings and communication into the same workflow. Supervisors can schedule sessions, record outcomes, "
    "and keep participants informed through notifications and linked communication channels.",
    size=11, after=6
)
add_picture(doc, "p1590_rId114.png", width=5.8,
            caption="The meeting calendar makes upcoming supervision sessions easy to see and manage.")

add_heading(doc, "Automatic reminders reduce missed follow-up", 2)
add_text(
    doc,
    "When a task or meeting is created or updated, the product can send a message through Discord. "
    "This connects formal tracking with the communication channel users already check.",
    size=10.7, after=5
)
add_picture(doc, "p1586_rId113.png", width=4.6,
            caption="The TaskFlow bot can notify a shared channel when work is assigned.")

doc.add_page_break()

# Page 5 — company value and next step
add_heading(doc, "Value for the company", 1)
add_feature_table(doc, [
    ("For management", "A reliable overview of activity, workload, completion, and areas requiring attention."),
    ("For supervisors", "Less time chasing updates and more time supporting people with the right guidance."),
    ("For learners", "Clear responsibilities, visible deadlines, and a consistent place to ask questions and show progress."),
])

add_picture(doc, "p1628_rId123.png", width=5.85,
            caption="Reports turn everyday activity into a simple management view of completion, delays, and project delivery.")

add_heading(doc, "Proposed engagement", 1)
add_text(
    doc,
    "The recommended next step is a short company pilot using one team or training group. "
    "During the pilot, the product can be configured for the company’s roles, terminology, workflow, and communication preferences.",
    size=10.8, after=6
)
add_process_step(doc, "1", "Discovery", "Confirm the company’s supervision process and the people who will use TaskFlow.")
add_process_step(doc, "2", "Pilot setup", "Prepare users, boards, permissions, reminders, and a small set of active projects.")
add_process_step(doc, "3", "Evaluation", "Collect feedback, review adoption and progress visibility, then agree on improvements.")
add_process_step(doc, "4", "Rollout decision", "Define the final scope, support model, timeline, and commercial terms.")

add_heading(doc, "Requested company input", 2)
add_text(
    doc,
    "A pilot sponsor, a small user group, access to the current workflow, and feedback from administrators, supervisors, and learners.",
    size=10.7, after=10
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font_run(p.add_run("TaskFlow is ready to demonstrate."), size=15, bold=True, color=NAVY)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.paragraph_format.space_after = Pt(0)
font_run(p2.add_run("Prepared by Reem Alhalwachi"), size=11, bold=True, color=TEAL)

doc.save(OUT)
print(OUT)
